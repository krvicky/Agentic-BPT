"""Local parsers preserve text/table blocks and make unsupported content explicit."""
from io import BytesIO
import hashlib
from pathlib import Path
import zipfile
from typing import Callable
from .models import ParseIssue, ParsedDocument, SourceBlock, SourceDocument, stable_id
from ..tracing import contains_secret, redact

MAX_FILE_BYTES = 25 * 1024 * 1024
MAX_EXPANDED_BYTES = 100 * 1024 * 1024

def block_id(document_id: str, kind: str, locator: str, text: str, rows: list[list[str]],
             page: int | None = None, slide: int | None = None, section: str = "") -> str:
    return stable_id("block", document_id, kind, locator, text, rows, page, slide, section)

def parse_bytes(filename: str, data: bytes) -> ParsedDocument:
    name = str(redact(Path(filename).name))
    format = Path(filename).suffix.lower().lstrip(".")
    digest = hashlib.sha256(data).hexdigest()
    id = stable_id("doc", digest, format)
    issues: list[ParseIssue] = []
    blocks: list[SourceBlock] = []

    def issue(code: str, message: str, locator: str = "document", affects: bool = True) -> None:
        issues.append(ParseIssue(code, message, locator, affects))

    def add(kind: str, locator: str, text: str, *, page: int | None = None, slide: int | None = None,
            section: str = "", rows: list[list[str]] | None = None) -> None:
        rows = rows or []
        if text.strip():
            blocks.append(SourceBlock(block_id(id, kind, locator, text, rows, page, slide, section), id, kind, locator, text, page, slide, section, rows))

    if len(data) > MAX_FILE_BYTES:
        issue("file_size_limit", "Document exceeds the configured 25 MiB parsing limit")
    elif format not in ("pdf", "pptx", "docx"):
        issue("unsupported_format", "Supported formats are PDF, PPTX and DOCX")
    else:
        try:
            if format in ("pptx", "docx"):
                with zipfile.ZipFile(BytesIO(data)) as archive:
                    if sum(info.file_size for info in archive.infolist()) > MAX_EXPANDED_BYTES:
                        raise ValueError("Expanded archive exceeds the parsing limit")
                    special = ("embeddings/", "diagrams/", "vbaProject", "footnotes.xml", "endnotes.xml", "comments.xml")
                    for member in archive.namelist():
                        if any(part in member for part in special):
                            issue("unsupported_part", "Embedded, diagram, revision or annotation content is retained in the original but not fully parsed", member)
            if format == "pdf":
                _pdf(data, add, issue)
            elif format == "pptx":
                _pptx(data, add, issue)
            else:
                _docx(data, add, issue)
        except Exception as exc:
            issue("document_parse_error", "Parser failed (" + type(exc).__name__ + "); extracted blocks, if any, are retained")
    archive_name = "sources/" + id + "." + (format if format in ("pdf", "pptx", "docx") else "bin")
    if contains_secret(blocks):
        blocks.clear()
        archive_name = ""
        issue("sensitive_source", "Recognizable credentials were detected; original content and parsed text were not archived")
    if not blocks:
        issue("no_extractable_text", "No usable text was extracted; this does not establish that the document contains no process information")
    status = "failed" if not blocks else "partial" if any(i.affects_coverage for i in issues) else "parsed"
    document = SourceDocument(id, digest, name, format, archive_name, status, [b.id for b in blocks], issues)
    return ParsedDocument(document, blocks)

def parse_document(path: str | Path) -> tuple[ParsedDocument, bytes | None]:
    source = Path(path)
    try:
        with source.open("rb") as handle:
            data = handle.read(MAX_FILE_BYTES + 1)
        if len(data) > MAX_FILE_BYTES:
            # Do not archive a truncated source as if it were the original.
            raise OverflowError("File size limit")
        return parse_bytes(source.name, data), data
    except OSError as exc:
        code = type(exc).__name__
    except OverflowError:
        code = "FileSizeLimit"
    name = str(redact(source.name))
    document = SourceDocument(stable_id("unreadable", name), "", name, source.suffix.lower().lstrip("."), "", "failed", [],
                              [ParseIssue("file_read_error", "Source could not be read (" + code + ")", "document")])
    return ParsedDocument(document, []), None

def _pdf(data: bytes, add: Callable[..., None], issue: Callable[..., None]) -> None:
    import pdfplumber
    with pdfplumber.open(BytesIO(data)) as pdf:
        for number, page in enumerate(pdf.pages, 1):
            locator = f"page:{number}"
            try:
                tables = []
                try:
                    tables = page.find_tables()
                except Exception:
                    issue("table_parse_error", "Table detection failed; page text is still attempted", locator)
                extracted_tables = []
                for index, table in enumerate(tables, 1):
                    try:
                        rows = [[cell or "" for cell in row] for row in table.extract()]
                        text = "\n".join("\t".join(row) for row in rows)
                        if text.strip():
                            extracted_tables.append((index, table.bbox, rows, text))
                    except Exception:
                        issue("table_parse_error", "Detected table parsing failed; its raw page text is retained", locator + f"/table:{index}")
                # Exclude only successfully extracted tables, so failures never lose text.
                boxes = [box for _, box, _, _ in extracted_tables]
                def outside_tables(obj):
                    if obj.get("object_type") != "char":
                        return True
                    return not any(obj["x0"] >= x0 and obj["x1"] <= x1 and obj["top"] >= top and obj["bottom"] <= bottom
                                   for x0, top, x1, bottom in boxes)
                text = page.filter(outside_tables).extract_text() if boxes else page.extract_text()
                add("text", locator + "/text", text or "", page=number, section="Page " + str(number))
                for index, box, rows, table_text in extracted_tables:
                    add("table", locator + f"/table:{index}", table_text, page=number, section="Page " + str(number), rows=rows)
                if page.images:
                    issue("image_content", "Images/scans are not OCR-processed; image information may be missing", locator)
                if page.curves or any(not any(
                        obj["x0"] >= x0 and obj["x1"] <= x1 and obj["top"] >= top and obj["bottom"] <= bottom
                        for x0, top, x1, bottom in boxes) for obj in page.rects + page.lines):
                    issue("vector_content", "Vector diagrams and visual relationships are not interpreted", locator)
                if page.annots:
                    issue("annotations", "PDF annotations/forms are not interpreted as process evidence", locator)
                if not text and not tables:
                    issue("no_page_text", "Page has no extractable text; blank/scanned content remains uncertain", locator)
            except Exception as exc:
                issue("page_parse_error", "Page parsing failed (" + type(exc).__name__ + "); other pages continue", locator)

def _pptx(data: bytes, add: Callable[..., None], issue: Callable[..., None]) -> None:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    deck = Presentation(BytesIO(data))
    for number, slide in enumerate(deck.slides, 1):
        section = slide.shapes.title.text if slide.shapes.title and slide.shapes.title.has_text_frame else "Slide " + str(number)
        found = False
        def shapes(items, prefix):
            nonlocal found
            for index, shape in enumerate(items, 1):
                locator = prefix + f"/shape:{index}"
                try:
                    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                        shapes(shape.shapes, locator)
                    elif shape.has_table:
                        rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                        add("table", locator, "\n".join("\t".join(row) for row in rows), slide=number, section=section, rows=rows)
                        found = True
                        if any(cell.is_merge_origin or cell.is_spanned for row in shape.table.rows for cell in row.cells):
                            issue("merged_table", "Merged cell layout may need visual review", locator)
                    elif shape.has_text_frame:
                        add("text", locator, shape.text_frame.text, slide=number, section=section)
                        found = found or bool(shape.text_frame.text.strip())
                        if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
                            issue("visual_relationships", "Shape text is retained; diagram geometry is not interpreted", locator)
                    elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                        issue("image_content", "Slide image content is not OCR-processed", locator)
                    elif shape.has_chart:
                        issue("chart_content", "Chart data/relationships are not interpreted", locator)
                    elif shape.shape_type == MSO_SHAPE_TYPE.LINE:
                        issue("visual_relationships", "Connector relationships are not interpreted", locator)
                    else:
                        issue("unsupported_shape", "Non-text slide object is not interpreted", locator)
                except Exception as exc:
                    issue("shape_parse_error", "Shape parsing failed (" + type(exc).__name__ + "); other shapes continue", locator)
        shapes(slide.shapes, f"slide:{number}")
        try:
            if slide.has_notes_slide:
                frame = slide.notes_slide.notes_text_frame
                if frame is not None:
                    add("notes", f"slide:{number}/notes", frame.text, slide=number, section=section)
        except Exception:
            issue("notes_parse_error", "Speaker notes could not be parsed", f"slide:{number}/notes")
        if not found:
            issue("no_slide_text", "No main-slide text was extracted; image-only information may be missing", f"slide:{number}")

def _docx(data: bytes, add: Callable[..., None], issue: Callable[..., None]) -> None:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    document = Document(BytesIO(data))
    section = "Document body"
    issue("pagination_unavailable", "DOCX physical pages require layout rendering; citations use headings and block locators", affects=False)

    def content(container, prefix, default_kind="text"):
        nonlocal section
        for index, item in enumerate(container.iter_inner_content(), 1):
            locator = prefix + f"/block:{index}"
            try:
                if isinstance(item, Paragraph):
                    if item.style is not None and item.style.name.startswith("Heading") and item.text.strip():
                        section = item.text
                    add(default_kind, locator, item.text, section=section)
                    for tag, code in [("w:drawing", "image_content"), ("w:txbxContent", "textbox_content"), ("w:ins", "tracked_changes"), ("w:del", "tracked_changes"), ("w:object", "embedded_content")]:
                        if item._p.xpath(".//" + tag):
                            issue(code, "Drawing/textbox/revision/embedded content is retained in the source but may not be represented by paragraph text", locator)
                elif isinstance(item, Table):
                    rows = [[cell.text for cell in row.cells] for row in item.rows]
                    add("table", locator, "\n".join("\t".join(row) for row in rows), section=section, rows=rows)
                    if item._tbl.xpath(".//w:gridSpan") or item._tbl.xpath(".//w:vMerge"):
                        issue("merged_table", "Merged table layout may need visual review", locator)
                    if any(item._tbl.xpath(".//" + tag) for tag in ("w:drawing", "w:txbxContent", "w:ins", "w:del", "w:object")):
                        issue("unsupported_table_content", "Table images/textboxes/revisions/embedded content are not fully interpreted", locator)
                    seen_cells = set()
                    for r, row in enumerate(item.rows, 1):
                        for c, cell in enumerate(row.cells, 1):
                            if cell._tc in seen_cells:
                                continue
                            seen_cells.add(cell._tc)
                            for t, nested in enumerate(cell.tables, 1):
                                # Nested tables are emitted separately, not silently lost.
                                rows = [[inner.text for inner in row.cells] for row in nested.rows]
                                add("table", locator + f"/row:{r}/cell:{c}/table:{t}",
                                    "\n".join("\t".join(row) for row in rows), section=section, rows=rows)
                                if any(inner.tables for row in nested.rows for inner in row.cells):
                                    issue("deep_nested_table", "Table nesting deeper than one level is not fully parsed", locator)
            except Exception as exc:
                issue("block_parse_error", "DOCX block parsing failed (" + type(exc).__name__ + "); other blocks continue", locator)
    content(document, "body")
    seen_parts = set()
    for number, doc_section in enumerate(document.sections, 1):
        for kind, containers in [("header", [doc_section.header, doc_section.first_page_header, doc_section.even_page_header]),
                                 ("footer", [doc_section.footer, doc_section.first_page_footer, doc_section.even_page_footer])]:
            for variant, container in enumerate(containers, 1):
                try:
                    if not container._has_definition or container.part.partname in seen_parts:
                        continue
                    seen_parts.add(container.part.partname)
                    section = f"Section {number} {kind}"
                    content(container, f"section:{number}/{kind}:{variant}", kind)
                except Exception:
                    issue("header_footer_error", "Header or footer parsing failed", f"section:{number}/{kind}")
    # Content controls and revisions outside ordinary paragraphs can be omitted by the library.
    for tag, code in [("w:sdt", "content_control"), ("w:ins", "tracked_changes"), ("w:del", "tracked_changes")]:
        if document.element.xpath(".//" + tag):
            issue(code, "Special DOCX content outside normal blocks may be incomplete")
