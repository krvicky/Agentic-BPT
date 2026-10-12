"""Synthetic binary documents; no confidential files or paid API access."""
from io import BytesIO
from pathlib import Path
import pytest

@pytest.fixture
def sparse_slide(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Quotation process"
    box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(3))
    box.text_frame.text = "Sales collects vendor rates.\nSales compiles a branded quote.\nSales sends the quote to the customer.\nTo-Be: automate quotation preparation."
    path = tmp_path / "sparse.pptx"
    deck.save(path)
    return path

@pytest.fixture
def detailed_sop(tmp_path):
    from docx import Document
    doc = Document()
    doc.add_heading("Quotation SOP", 1)
    doc.add_paragraph("Sales collects vendor rates.")
    doc.add_paragraph("If a rate is unavailable, Sales escalates to Procurement.")
    doc.add_paragraph("The quotation must be sent within two working days.")
    doc.add_paragraph("Manual re-entry causes delays.")
    doc.add_heading("Controls", 2)
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "System"
    table.cell(0, 1).text = "Document"
    table.cell(1, 0).text = "CRM"
    table.cell(1, 1).text = "Branded quotation"
    doc.sections[0].header.paragraphs[0].text = "Controlled SOP"
    path = tmp_path / "detailed.docx"
    doc.save(path)
    return path

@pytest.fixture
def conflicting_documents(tmp_path):
    from docx import Document
    paths = []
    for filename, value in [("policy-a.docx", "10"), ("policy-b.docx", "20")]:
        doc = Document()
        doc.add_paragraph("Sales applies a " + value + " percent markup.")
        path = tmp_path / filename
        doc.save(path)
        paths.append(path)
    return paths

@pytest.fixture
def text_pdf(tmp_path):
    from reportlab.pdfgen.canvas import Canvas
    path = tmp_path / "process.pdf"
    canvas = Canvas(str(path))
    canvas.drawString(72, 750, "Sales sends the quote to the customer.")
    canvas.showPage()
    canvas.save()
    return path
