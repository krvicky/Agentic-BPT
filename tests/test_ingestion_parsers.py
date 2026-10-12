from copy import deepcopy
from dataclasses import asdict
from io import BytesIO
import pytest
from bpt_process_agent.ingestion.parsers import parse_document, parse_bytes
from bpt_process_agent.ingestion.models import Registry, Citation, CandidateFact
from bpt_process_agent.ingestion.validation import validate_candidate, extraction_candidates, EvidenceValidationError

def registry(parsed):
    return Registry(documents={parsed.document.id: parsed.document}, blocks={b.id: b for b in parsed.blocks})

def test_sparse_slide_preserves_exact_text_and_slide(sparse_slide):
    parsed, data = parse_document(sparse_slide)
    assert parsed.document.status == "parsed"
    assert data == sparse_slide.read_bytes()
    assert all(b.slide == 1 and b.page is None for b in parsed.blocks)
    assert any("To-Be: automate" in b.text for b in parsed.blocks)
    assert all(b.section == "Quotation process" for b in parsed.blocks)

def test_sop_tables_headings_and_headers(detailed_sop):
    parsed, _ = parse_document(detailed_sop)
    assert parsed.document.status == "parsed"
    table = next(b for b in parsed.blocks if b.kind == "table")
    assert table.rows == [["System", "Document"], ["CRM", "Branded quotation"]]
    assert table.section == "Controls"
    assert any(b.kind == "header" and b.text == "Controlled SOP" for b in parsed.blocks)
    assert all(b.page is None for b in parsed.blocks)
    assert any(i.code == "pagination_unavailable" and not i.affects_coverage for i in parsed.document.issues)

def test_pdf_page_provenance(text_pdf):
    parsed, _ = parse_document(text_pdf)
    assert parsed.document.status == "parsed"
    assert parsed.blocks[0].page == 1
    assert "Sales sends" in parsed.blocks[0].text

@pytest.mark.parametrize("filename,data,code", [
    ("broken.pdf", b"bad", "document_parse_error"),
    ("broken.docx", b"bad", "document_parse_error"),
    ("broken.pptx", b"bad", "document_parse_error"),
    ("source.xlsx", b"bad", "unsupported_format"),
])
def test_failures_explicit(filename, data, code):
    parsed = parse_bytes(filename, data)
    assert parsed.document.status == "failed"
    assert any(i.code == code for i in parsed.document.issues)
    assert any(i.code == "no_extractable_text" for i in parsed.document.issues)

def test_missing_file_is_recorded(tmp_path):
    parsed, data = parse_document(tmp_path / "missing.docx")
    assert parsed.document.status == "failed"
    assert data is None and not parsed.document.archive_name

def test_blank_pdf_never_claims_no_process_information(tmp_path):
    from reportlab.pdfgen.canvas import Canvas
    path = tmp_path / "blank.pdf"
    c = Canvas(str(path)); c.showPage(); c.save()
    parsed, _ = parse_document(path)
    assert parsed.document.status == "failed"
    assert any(i.code == "no_page_text" for i in parsed.document.issues)

def test_image_content_flagged_even_when_text_present(tmp_path):
    from docx import Document
    from PIL import Image
    image = tmp_path / "scan.png"
    Image.new("RGB", (10, 10)).save(image)
    document = Document()
    document.add_paragraph("Sales sends quotes.")
    document.add_picture(str(image))
    path = tmp_path / "image.docx"
    document.save(path)
    parsed, _ = parse_document(path)
    assert parsed.document.status == "partial"
    assert any(i.code == "image_content" for i in parsed.document.issues)

def test_scanned_pdf_flagged(tmp_path):
    from reportlab.pdfgen.canvas import Canvas
    from PIL import Image
    image = tmp_path / "scan.png"
    Image.new("RGB", (20, 20)).save(image)
    path = tmp_path / "scan.pdf"
    c = Canvas(str(path)); c.drawImage(str(image), 50, 50); c.save()
    parsed, _ = parse_document(path)
    assert parsed.document.status == "failed"
    assert any(i.code == "image_content" for i in parsed.document.issues)

def test_partial_page_failure_keeps_other_page(monkeypatch, text_pdf):
    import pdfplumber
    from bpt_process_agent.ingestion import parsers
    class Broken:
        def find_tables(self): raise ValueError()
        def extract_text(self): raise ValueError()
    real = pdfplumber.open(text_pdf)
    good = real.pages[0]
    class Fake:
        pages = [Broken(), good]
        def __enter__(self): return self
        def __exit__(self, *args): pass
    monkeypatch.setattr(pdfplumber, "open", lambda *a, **k: Fake())
    parsed = parse_bytes("partial.pdf", text_pdf.read_bytes())
    assert parsed.document.status == "partial"
    assert any(b.page == 2 for b in parsed.blocks)
    assert any(i.code == "page_parse_error" for i in parsed.document.issues)
    real.close()

def test_secret_source_not_archived(tmp_path):
    from docx import Document
    d = Document(); d.add_paragraph("OPENROUTER_API_KEY=sk-abcdefghijklmno")
    path = tmp_path / "secret.docx"; d.save(path)
    parsed, _ = parse_document(path)
    assert not parsed.blocks and not parsed.document.archive_name
    assert any(i.code == "sensitive_source" for i in parsed.document.issues)

@pytest.mark.parametrize("change", ["offset", "quote", "block", "ambiguous", "tags"])
def test_citation_and_classification_negative_cases(change, detailed_sop):
    parsed, _ = parse_document(detailed_sop)
    block = next(b for b in parsed.blocks if b.text == "Sales collects vendor rates.")
    fact = CandidateFact(block.text, ["activity", "role"], "as_is", [Citation(block.id, 0, len(block.text), block.text)], "")
    if change == "offset": fact.citations[0].end += 1
    if change == "quote": fact.citations[0].quote = "Sales approves quotes."
    if change == "block": fact.citations[0].block_id = "invented"
    if change == "ambiguous": fact.classification = "ambiguous"
    if change == "tags": fact.tags = ["activity", "activity"]
    with pytest.raises(EvidenceValidationError):
        validate_candidate(fact, registry(parsed))

def test_multitag_as_is_to_be_and_ambiguity_are_distinct(sparse_slide):
    parsed, _ = parse_document(sparse_slide)
    block = next(b for b in parsed.blocks if "To-Be" in b.text)
    quote = "To-Be: automate quotation preparation."
    start = block.text.index(quote)
    fact = CandidateFact(quote, ["activity", "system"], "to_be", [Citation(block.id, start, start + len(quote), quote)], "")
    assert extraction_candidates({"facts": [asdict(fact)]}, registry(parsed), {block.id}) == [fact]
    fact.classification = "ambiguous"; fact.ambiguity = "Timing is unspecified."
    validate_candidate(fact, registry(parsed))

def test_pdf_table_failure_retains_raw_text(monkeypatch, text_pdf):
    import pdfplumber
    from bpt_process_agent.ingestion.parsers import parse_bytes
    class BrokenTable:
        bbox = (0, 0, 500, 500)
        def extract(self): raise ValueError("table error")
    class Page:
        images, curves, annots, rects, lines = [], [], [], [], []
        def find_tables(self): return [BrokenTable()]
        def extract_text(self): return "Sales sends the quote to the customer."
        def filter(self, fn): raise AssertionError("Failed tables must not exclude text")
    class Fake:
        pages = [Page()]
        def __enter__(self): return self
        def __exit__(self, *args): pass
    monkeypatch.setattr(pdfplumber, "open", lambda *a, **k: Fake())
    parsed = parse_bytes("tables.pdf", text_pdf.read_bytes())
    assert parsed.document.status == "partial"
    assert parsed.blocks[0].text == "Sales sends the quote to the customer."
    assert any(i.code == "table_parse_error" for i in parsed.document.issues)

def test_nested_docx_table_preserved(tmp_path):
    from docx import Document
    d = Document(); outer = d.add_table(rows=1, cols=1)
    outer.cell(0, 0).text = "Sales rules"
    inner = outer.cell(0, 0).add_table(rows=1, cols=2)
    inner.cell(0, 0).text = "Exception"
    inner.cell(0, 1).text = "Sales escalates"
    path = tmp_path / "nested.docx"; d.save(path)
    parsed, _ = parse_document(path)
    assert any(b.rows == [["Exception", "Sales escalates"]] for b in parsed.blocks)

def test_slide_table_notes_and_scan_limitations(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches
    from PIL import Image
    deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Controls"
    table = slide.shapes.add_table(1, 2, Inches(1), Inches(1), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "Owner"; table.cell(0, 1).text = "Sales"
    slide.notes_slide.notes_text_frame.text = "Sales sends quotes."
    image = tmp_path / "image.png"; Image.new("RGB", (10, 10)).save(image)
    slide.shapes.add_picture(str(image), Inches(1), Inches(3))
    path = tmp_path / "objects.pptx"; deck.save(path)
    parsed, _ = parse_document(path)
    assert parsed.document.status == "partial"
    assert any(b.kind == "notes" and "Sales sends quotes." in b.text for b in parsed.blocks)
    assert any(b.rows == [["Owner", "Sales"]] for b in parsed.blocks)
    assert any(i.code == "image_content" for i in parsed.document.issues)

def test_expanded_archive_limit_is_explicit(monkeypatch, detailed_sop):
    from bpt_process_agent.ingestion import parsers
    monkeypatch.setattr(parsers, "MAX_EXPANDED_BYTES", 1)
    parsed, _ = parsers.parse_document(detailed_sop)
    assert parsed.document.status == "failed"
    assert any(i.code == "document_parse_error" for i in parsed.document.issues)

def test_exact_unicode_offsets(tmp_path):
    from docx import Document
    d = Document(); d.add_paragraph("Ventes prépare un devis — délai: ٢ jours.")
    path = tmp_path / "unicode.docx"; d.save(path)
    parsed, _ = parse_document(path)
    block = parsed.blocks[0]
    quote = "prépare un devis"
    start = block.text.index(quote)
    fact = CandidateFact("Ventes prépare un devis.", ["activity", "role"], "as_is",
                         [Citation(block.id, start, start + len(quote), quote)], "")
    validate_candidate(fact, registry(parsed))
